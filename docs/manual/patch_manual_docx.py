"""Post-render docx patch for the viewer manual (run by build_manual.ps1 after `quarto render`).

Applies what the reference.docx style alone doesn't:
  - center every table horizontally,
  - center every figure paragraph (image) and its caption,
  - give every body prose paragraph a first-line indent (0.25"),
  - clear the build machine from the file: Quarto writes the absolute path of its callout icons
    into their alt text (``descr``), and the core properties may name an author or editor.
Headings, list items, TOC entries, captions, images and empty paragraphs are left un-indented.

    python patch_manual_docx.py <a.docx> [<b.docx> ...]
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

INDENT = Inches(0.25)
SKIP_PREFIX = ("Heading", "Title", "Subtitle", "TOC", "Author", "Date")


def _has_drawing(p):
    return "w:drawing" in p._p.xml


def _is_list(p):
    pPr = p._p.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:numPr")) is not None


#: An alt text that is an absolute path (a drive letter, a UNC share or a POSIX root).
_ABSOLUTE = re.compile(r"^(?:[A-Za-z]:[\\/]|\\\\|/)")


def _scrub(doc):
    """Blank every absolute-path alt text and the personal core properties; count the alt texts."""
    cleared = 0
    body = doc.element.body
    for tag in (qn("pic:cNvPr"), qn("wp:docPr")):
        for el in body.iter(tag):
            if _ABSOLUTE.match(el.get("descr") or ""):
                el.set("descr", "")
                cleared += 1
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    return cleared


def patch(path):
    doc = Document(path)
    n_descr = _scrub(doc)
    for t in doc.tables:                                  # #1: center tables
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
    n_ind = 0
    for p in doc.paragraphs:
        style = p.style.name or ""
        if _has_drawing(p):                               # #1: center figure image
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if "Caption" in style:                            # center caption under the figure
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if any(style.startswith(s) for s in SKIP_PREFIX):
            continue
        if _is_list(p) or not p.text.strip():
            continue
        p.paragraph_format.first_line_indent = INDENT     # #4: body prose first-line indent
        n_ind += 1
    doc.save(path)
    print(f"patched {path}  (tables={len(doc.tables)}, indented_paras={n_ind}, cleared_alt_paths={n_descr})")


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        patch(arg)
